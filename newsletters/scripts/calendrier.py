"""Génère newsletters/calendrier-editorial.xlsx (octobre 2026 à décembre 2027).

Règles d'envoi :
- B2B : le 15. Samedi -> vendredi 14, dimanche -> lundi 16, jour férié -> jour ouvré suivant.
- Trucker et Owner : le dernier vendredi du mois. Férié, 24 ou 31 décembre -> vendredi précédent.
- Nouveautés Trucker : le lundi qui suit l'envoi Trucker. Férié -> mardi.
"""
from datetime import date, timedelta
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

SORTIE = Path(__file__).resolve().parent.parent / "calendrier-editorial.xlsx"

FERIES = {  # France et Italie (langues Trucker et Owner)
    date(2026, 11, 1): "Toussaint", date(2026, 11, 11): "Armistice (FR)",
    date(2026, 12, 8): "Immacolata (IT)", date(2026, 12, 25): "Noël", date(2026, 12, 26): "Santo Stefano (IT)",
    date(2027, 1, 1): "Jour de l'an", date(2027, 1, 6): "Epifania (IT)",
    date(2027, 3, 28): "Pâques", date(2027, 3, 29): "Lundi de Pâques",
    date(2027, 4, 25): "Liberazione (IT)", date(2027, 5, 1): "Fête du travail",
    date(2027, 5, 6): "Ascension (FR)", date(2027, 5, 8): "Victoire 1945 (FR)",
    date(2027, 5, 17): "Lundi de Pentecôte (FR)", date(2027, 6, 2): "Festa della Repubblica (IT)",
    date(2027, 7, 14): "Fête nationale (FR)", date(2027, 8, 15): "Assomption",
    date(2027, 11, 1): "Toussaint", date(2027, 11, 11): "Armistice (FR)",
    date(2027, 12, 8): "Immacolata (IT)", date(2027, 12, 25): "Noël", date(2027, 12, 26): "Santo Stefano (IT)",
}
A_EVITER = {date(2026, 12, 24), date(2026, 12, 31), date(2027, 12, 24), date(2027, 12, 31)}

MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
        "août", "septembre", "octobre", "novembre", "décembre"]
JOURS = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]

TYPES = {
    "B2B": ("Partenaires et clients pros", "FR, EN, IT, ES, DE, NL, PL", "DCE6F7"),
    "Trucker": ("Conducteurs (app)", "FR, IT", "FFF7B3"),
    "Owner": ("Établissements sur Truckfly", "FR", "DFF0E0"),
    "Nouveautés Trucker": ("Conducteurs (app)", "FR, IT", "FBE3D0"),
}


def fr(d):
    return f"{JOURS[d.weekday()]} {d.day} {MOIS[d.month - 1]} {d.year}"


def date_b2b(a, m):
    d = date(a, m, 15)
    note = []
    if d.weekday() == 5:
        d, note = d - timedelta(1), ["Le 15 tombe un samedi : avancé au vendredi."]
    elif d.weekday() == 6:
        d, note = d + timedelta(1), ["Le 15 tombe un dimanche : décalé au lundi."]
    while d in FERIES or d.weekday() >= 5:
        note.append(f"{FERIES.get(d, 'Week-end')} : décalé au jour ouvré suivant.")
        d += timedelta(1)
    if m == 8:
        note.append("Creux de l'été : envisager une édition plus courte.")
    return d, " ".join(note)


def date_trucker(a, m):
    d = date(a + (m == 12), m % 12 + 1, 1) - timedelta(1)
    while d.weekday() != 4:
        d -= timedelta(1)
    note = []
    while d in FERIES or d in A_EVITER:
        raison = FERIES.get(d) or ("Veille de Noël" if d.day == 24 else "Saint-Sylvestre")
        note.append(f"{fr(d).capitalize()} : {raison}, avancé au vendredi précédent.")
        d -= timedelta(7)
    if d == date(2027, 3, 26):
        note.append("Vendredi saint (férié en Allemagne et en Espagne, pas en France ni en Italie).")
    return d, " ".join(note)


def date_nouveautes(vendredi):
    d = vendredi + timedelta(3)
    note = ""
    if d in FERIES:
        note = f"{FERIES[d]} : décalé au mardi."
        d += timedelta(1)
    return d, note


def lignes():
    out = []
    a, m = 2026, 10
    while (a, m) <= (2027, 12):
        cle = f"{MOIS[m - 1].capitalize()} {a}"
        b, nb = date_b2b(a, m)
        t, nt = date_trucker(a, m)
        n, nn = date_nouveautes(t)
        out += [(b, cle, "B2B", nb), (t, cle, "Trucker", nt), (t, cle, "Owner", nt), (n, cle, "Nouveautés Trucker", nn)]
        a, m = (a + 1, 1) if m == 12 else (a, m + 1)
    return sorted(out, key=lambda r: (r[0], list(TYPES).index(r[2])))


POLICE = "Arial"
BLEU = "061866"
fin = Side(style="thin", color="D0D4E0")
BORD = Border(left=fin, right=fin, top=fin, bottom=fin)
SAISIE = PatternFill("solid", fgColor="FFFBE6")


def entete(ws, ligne, titres, largeurs):
    for i, (t, l) in enumerate(zip(titres, largeurs), 1):
        c = ws.cell(ligne, i, t)
        c.font = Font(name=POLICE, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLEU)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORD
        ws.column_dimensions[c.column_letter].width = l
    ws.row_dimensions[ligne].height = 32


def main():
    wb = Workbook()
    cal = wb.active
    cal.title = "Calendrier"
    regles = wb.create_sheet("Règles")
    vue = wb.create_sheet("Vue par mois")

    # ---------- Règles et jours fériés ----------
    regles.column_dimensions["A"].width = 26
    regles.column_dimensions["B"].width = 34
    regles.column_dimensions["C"].width = 70
    regles["A1"] = "Mode d'emploi"
    regles["A1"].font = Font(name=POLICE, bold=True, size=14, color=BLEU)
    texte = [
        ("Cases à remplir", "Fond jaune clair dans Calendrier : Sujet principal, Objet, Statut, Responsable."),
        ("Exemple", "La ligne Trucker du 30 octobre 2026 est remplie."),
        ("B2B", "Le 15 du mois. Samedi : vendredi 14. Dimanche : lundi 16. Jour férié : jour ouvré suivant."),
        ("Trucker et Owner", "Le dernier vendredi du mois. Jour férié, 24 ou 31 décembre : vendredi précédent."),
        ("Nouveautés Trucker", "Le lundi qui suit l'envoi Trucker, seulement s'il y a une nouveauté. Jour férié : mardi."),
        ("Échéances (proposition)", "Texte FR validé : 5 jours ouvrés avant l'envoi. Traductions prêtes : 3 jours ouvrés avant. "
                                    "Programmé dans Brevo : 1 jour ouvré avant. Calculées à partir de la date d'envoi et des jours fériés ci-dessous."),
        ("Changer une date", "Modifier la date d'envoi dans Calendrier : jour, échéances et Vue par mois se mettent à jour."),
    ]
    for i, (k, v) in enumerate(texte, 3):
        regles.cell(i, 1, k).font = Font(name=POLICE, bold=True)
        c = regles.cell(i, 2, v)
        c.font = Font(name=POLICE)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        regles.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
        regles.row_dimensions[i].height = 30
    l0 = len(texte) + 5
    regles.cell(l0 - 1, 1, "Jours fériés pris en compte (France et Italie)").font = Font(name=POLICE, bold=True, size=12, color=BLEU)
    entete(regles, l0, ["Date", "Jour", "Férié"], [26, 34, 70])
    for i, (d, nom) in enumerate(sorted(FERIES.items()), l0 + 1):
        regles.cell(i, 1, d).number_format = "dd/mm/yyyy"
        regles.cell(i, 2, fr(d))
        regles.cell(i, 3, nom)
        for c in regles[i]:
            c.font = Font(name=POLICE)
            c.border = BORD
    plage_feries = f"'Règles'!$A${l0 + 1}:$A${l0 + len(FERIES)}"

    # ---------- Calendrier ----------
    titres = ["Date d'envoi", "Jour", "Mois", "Newsletter", "Cible", "Langues", "Sujet principal",
              "Objet de l'e-mail", "Statut", "Responsable", "Texte FR validé", "Traductions prêtes",
              "Programmé dans Brevo", "Note"]
    entete(cal, 1, titres, [13, 11, 15, 19, 26, 26, 44, 36, 14, 15, 14, 14, 14, 52])
    cal.freeze_panes = "E2"
    rows = lignes()
    for r, (d, cle, typ, note) in enumerate(rows, 2):
        cible, langues, couleur = TYPES[typ]
        sujet, objet = "À définir", ""
        statut = "Si nouveauté" if typ == "Nouveautés Trucker" else "À rédiger"
        if (d, typ) == (date(2026, 10, 30), "Trucker"):
            sujet = ("Heure d'hiver : trouver sa place avant la nuit. En bref : Profil Trucker, "
                     "amis et recommandations, Loi Montagne / obbligo invernale.")
            objet = "Depuis dimanche, la nuit te rattrape"
            statut = "Rédigé"
        valeurs = [d, f'=CHOOSE(WEEKDAY(A{r},2),"lundi","mardi","mercredi","jeudi","vendredi","samedi","dimanche")',
                   cle, typ, cible, langues, sujet, objet, statut, "",
                   f"=WORKDAY(A{r},-5,{plage_feries})",
                   f'=IF(F{r}="FR","—",WORKDAY(A{r},-3,{plage_feries}))',
                   f"=WORKDAY(A{r},-1,{plage_feries})", note]
        for col, v in enumerate(valeurs, 1):
            c = cal.cell(r, col, v)
            c.font = Font(name=POLICE, bold=(col == 4))
            c.border = BORD
            c.alignment = Alignment(vertical="top", wrap_text=col in (7, 8, 14))
            if col in (1, 11, 12, 13):
                c.number_format = "dd/mm/yyyy"
                c.alignment = Alignment(vertical="top", horizontal="left")
            if col == 4:
                c.fill = PatternFill("solid", fgColor=couleur)
            if col in (7, 8, 9, 10):
                c.fill = SAISIE
    der = len(rows) + 1
    cal.auto_filter.ref = f"A1:N{der}"
    dv = DataValidation(type="list", formula1='"À rédiger,Si nouveauté,En rédaction,Rédigé,En traduction,Validé,Programmé,Envoyé,Annulé"', allow_blank=True)
    cal.add_data_validation(dv)
    dv.add(f"I2:I{der}")
    gris = Font(name=POLICE, color="8A8F9C")
    cal.conditional_formatting.add(f"A2:N{der}", FormulaRule(formula=['OR($I2="Envoyé",$I2="Annulé")'], font=gris))
    cal.conditional_formatting.add(f"I2:I{der}", FormulaRule(formula=['$I2="Programmé"'], fill=PatternFill("solid", fgColor="C8E6C9")))

    # ---------- Vue par mois ----------
    entete(vue, 1, ["Mois", "B2B", "Trucker", "Owner", "Nouveautés Trucker"], [18, 30, 30, 30, 30])
    mois_liste = list(dict.fromkeys(r[1] for r in rows))
    for i, cle in enumerate(mois_liste, 2):
        vue.cell(i, 1, cle).font = Font(name=POLICE, bold=True)
        vue.cell(i, 1).border = BORD
        for j, typ in enumerate(TYPES, 2):
            f = (f'=_xlfn.MINIFS(Calendrier!$A$2:$A${der},Calendrier!$C$2:$C${der},$A{i},'
                 f'Calendrier!$D$2:$D${der},"{typ}")')
            c = vue.cell(i, j, f)
            c.number_format = "dddd dd/mm/yyyy"
            c.font = Font(name=POLICE)
            c.border = BORD
            c.alignment = Alignment(horizontal="left")
    vue.freeze_panes = "B2"

    for ws in (cal, vue, regles):
        ws.sheet_view.zoomScale = 110
    wb.save(SORTIE)
    print(SORTIE, len(rows), "envois")
    for d, cle, typ, note in rows:
        if note:
            print(" ", fr(d), typ, "|", note)


if __name__ == "__main__":
    main()
